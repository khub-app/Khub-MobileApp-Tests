"""Regenerates data/test_data.xlsx from the definitions below.

Mirrors khub-pos-tests' data/test_data.xlsx layout so the two suites are
driven the same way:

  * RunConfig   - which environment + which suite(s) run_mobile_tests.ps1 runs
  * Tags        - one row per test file; a "Yes" in a marker column makes
                  conftest.py auto-apply that pytest marker at collection time
  * TestConfigs - misc named test inputs the app UI never surfaces

The .xlsx itself is gitignored-friendly binary; edit THIS file and re-run
  venv\\Scripts\\python.exe scripts/gen_test_data_xlsx.py
to rebuild it, so the source of truth stays reviewable in git.
"""
import os

import openpyxl

_OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    "data", "test_data.xlsx")

RUN_CONFIG = [
    ("Setting", "Value", None),
    ("Environment", "emulator", "emulator or real_device - matches config/config.yaml"),
    (None, None, None),
    ("Suite", "Enabled", None),
    ("Smoke", "Yes", None),
    ("Regression", "No", None),
    (None, None, None),
    ("How to run", "Command", None),
    ("After editing this sheet / the Tags sheet, run:", ".\\run_mobile_tests.ps1", None),
    ("Run a specific suite regardless of this sheet:", "venv\\Scripts\\python.exe -m pytest tests/ -v -m regression --env=emulator", None),
    ("Run one file:", "venv\\Scripts\\python.exe -m pytest tests/fulfillment/test_online_sale_fulfillment.py -v --env=emulator", None),
]

TAGS = [
    ("FilePath", "smoke", "regression", "Date Added"),
    ("tests/login/test_login.py", "Yes", "Yes", "2026-09-03"),
    # Long multi-module E2E - regression only, not smoke.
    ("tests/fulfillment/test_online_sale_fulfillment.py", "No", "Yes", "2026-09-03"),
    ("tests/fulfillment/test_uom_multi_unit_sale_fulfillment.py", "No", "Yes", "2026-10-29"),
]

TEST_CONFIGS = [
    ("ConfigName", "Value", "Description"),
    ("dashboard_revenue_card", "Total Sale/Revenue",
     "content-desc of the Dashboard KPI card asserted after login"),
    ("online_sale_order_type", "Online",
     "Order Type used by the fulfillment E2E (see data/test_data.yaml online_sale)"),
    ("online_sale_payment_mode", "Cash",
     "Payment mode used to fully pay the fulfillment E2E's sale"),
]


def _write_sheet(ws, rows):
    for row in rows:
        ws.append(list(row))


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RunConfig"
    _write_sheet(ws, RUN_CONFIG)
    _write_sheet(wb.create_sheet("Tags"), TAGS)
    _write_sheet(wb.create_sheet("TestConfigs"), TEST_CONFIGS)
    wb.save(_OUT)
    print(f"[OK] wrote {_OUT}")


if __name__ == "__main__":
    main()
